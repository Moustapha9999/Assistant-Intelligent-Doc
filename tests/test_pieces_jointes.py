"""Pièces jointes multimodales : validation, extraction, transmission au modèle, persistance.

Tests hors ligne par défaut. `TESTS_LIVE=1` (avec GROQ_API_KEY) active les tests qui
appellent réellement le modèle de vision et le LLM.
"""

import base64
import io
import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import config
from app import gestionnaire_fichiers as gf
from app.historique_sqlite import HistoriqueSQLite
from securite.acces import AccesRefuse

if os.getenv("TESTS_LIVE") == "1":
    from dotenv import load_dotenv

    load_dotenv(Path(__file__).resolve().parents[1] / ".env")
LIVE = os.getenv("TESTS_LIVE") == "1" and bool(os.getenv("GROQ_API_KEY"))
live = pytest.mark.skipif(not LIVE, reason="TESTS_LIVE=1 et GROQ_API_KEY requis")


class _Upload(io.BytesIO):
    """Imite l'UploadedFile Streamlit (nom + octets)."""

    def __init__(self, nom, data):
        super().__init__(data)
        self.name = nom


# ── Fabriques de fichiers ───────────────────────────────────────────────────

def _png_texte(lignes, taille=(640, 200), rouge=True) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", taille, "white")
    d = ImageDraw.Draw(img)
    try:
        police = ImageFont.truetype("arial.ttf", 28)
    except OSError:
        police = ImageFont.load_default()
    for i, l in enumerate(lignes):
        d.text((20, 20 + i * 45), l, fill="black", font=police)
    if rouge:
        d.rectangle((taille[0] - 110, 30, taille[0] - 30, 110), fill="red")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _pdf_texte(pages) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    for texte in pages:
        pdf.add_page()
        pdf.multi_cell(0, 8, texte)
    return bytes(pdf.output())


def _pdf_scanne(lignes) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.image(io.BytesIO(_png_texte(lignes, taille=(1200, 400), rouge=False)), x=10, y=10, w=190)
    return bytes(pdf.output())


def _pdf_chiffre() -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_encryption(owner_password="proprio", user_password="secret")
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, "Confidentiel")
    return bytes(pdf.output())


def _docx() -> bytes:
    import docx

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Societe Alpha - Rapport interne"
    d.add_heading("Rapport annuel", level=1)
    d.add_paragraph("Le chiffre d'affaires a progressé de 12 % en 2025.")
    d.add_heading("Risques", level=2)
    d.add_paragraph("Retard fournisseur", style="List Bullet")
    d.add_paragraph("Hausse des coûts énergétiques", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Indicateur", "Valeur"
    t.cell(1, 0).text, t.cell(1, 1).text = "Marge", "18 %"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventes"
    ws.append(["Mois", "Region", "Montant"])
    for mois, region, montant in [("Janvier", "Nord", 1200), ("Février", "Sud", 800),
                                  ("Mars", "Nord", 1500), ("Avril", "Est", 9900)]:
        ws.append([mois, region, montant])
    ws["E1"], ws["E2"] = "Total", "=SUM(C2:C5)"
    wb.create_sheet("Notes").append(["Commentaire"])
    wb["Notes"].append(["Avril : commande exceptionnelle"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Feuille de route"
    s.placeholders[1].text = "Lancement V2 en mars"
    s.notes_slide.notes_text_frame.text = "Insister sur la sécurité"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class _FauxVision:
    """Doublure du client LLM : enregistre ce qui est réellement envoyé au modèle de vision."""

    def __init__(self, reponse="Texte lu : FACTURE 4271. Un carré rouge à droite."):
        self.appels = []
        self.reponse = reponse
        self.derniere_erreur_vision = None

    def analyser_image(self, image_b64, media_type, question):
        self.appels.append({"b64": image_b64, "media": media_type, "question": question})
        if self.reponse is None:
            self.derniere_erreur_vision = "le service d'analyse d'images est indisponible pour le moment."
            return "", 0
        return self.reponse, 42


# ── 1-2. Images : contenu visuel réellement transmis à la vision ────────────

def test_image_contenu_visuel_transmis_au_modele():
    png = _png_texte(["FACTURE 4271"])
    f = gf.traiter_fichier(_Upload("capture.png", png))
    assert f["erreur"] is None and f["type"] == "image"
    assert f["media_type"] in {"image/png", "image/jpeg"} and f["resume"].startswith("Image · 640×200")

    vision = _FauxVision()
    gf.preparer_contenus_multimodaux([f], question="Décris cette image", llm_client=vision)
    assert len(vision.appels) == 1
    from PIL import Image

    recu = Image.open(io.BytesIO(base64.b64decode(vision.appels[0]["b64"])))
    assert recu.size == (640, 200)  # l'image elle-même, pas un nom ni un chemin
    assert "Décris cette image" in vision.appels[0]["question"]
    assert f["vision_ok"] is True
    prompt = gf.formater_pour_prompt([f], question="Décris cette image")
    assert "Image analysée : capture.png" in prompt and "FACTURE 4271" in prompt


def test_capture_ecran_question_transmise_et_echec_vision_explicite():
    f = gf.traiter_fichier(_Upload("ecran.png", _png_texte(["Erreur 504"])))
    vision = _FauxVision(reponse=None)
    gf.preparer_contenus_multimodaux([f], question="Explique le message d'erreur", llm_client=vision)
    assert "Explique le message d'erreur" in vision.appels[0]["question"]
    assert f["vision_ok"] is False
    assert any("Image non analysée" in a for a in f["avertissements"])
    prompt = gf.formater_pour_prompt([f])
    assert "NON analysé" in prompt and "Erreur 504" not in prompt


def test_grandes_images_reduites_et_limite_par_message(monkeypatch):
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (5000, 3000), "blue").save(buf, "JPEG")
    f = gf.traiter_fichier(_Upload("grand.jpg", buf.getvalue()))
    assert max(Image.open(io.BytesIO(base64.b64decode(f["image_b64"]))).size) == 2000

    monkeypatch.setattr(config, "FICHIERS_MAX_IMAGES_VISION", 1)
    images = [gf.traiter_fichier(_Upload(f"i{i}.png", _png_texte(["x"]))) for i in range(2)]
    vision = _FauxVision()
    gf.preparer_contenus_multimodaux(images, llm_client=vision)
    assert len(vision.appels) == 1 and images[1]["vision_ok"] is False


# ── 3-4. PDF textuel et PDF numérisé (OCR) ──────────────────────────────────

def test_pdf_textuel_pages_referencees():
    pdf = _pdf_texte([
        "Rapport 2025. Introduction generale du projet Atlas.",
        "Chapitre 2. Le budget total s'eleve a 48 000 euros.",
        "Conclusion. Le projet est livre en avance.",
    ])
    f = gf.traiter_fichier(_Upload("rapport.pdf", pdf))
    assert f["erreur"] is None
    assert f["resume"] == "PDF · 3 pages"
    assert [s["ref"] for s in f["segments"]] == ["page 1", "page 2", "page 3"]
    prompt = gf.formater_pour_prompt([f], question="Résume ce rapport")
    assert "--- page 2 ---" in prompt and "48 000 euros" in prompt


def test_pdf_numerise_ocr_par_la_vision():
    f = gf.traiter_fichier(_Upload("scan.pdf", _pdf_scanne(["CONTRAT DE BAIL", "Loyer 950 euros"])))
    assert f["erreur"] is None
    assert len(f["pages_ocr"]) == 1 and any("numérisé" in a for a in f["avertissements"])
    vision = _FauxVision(reponse="CONTRAT DE BAIL\nLoyer 950 euros")
    gf.preparer_contenus_multimodaux([f], question="Quel est le loyer ?", llm_client=vision)
    assert len(vision.appels) == 1 and "Transcris" in vision.appels[0]["question"]
    assert "pages_ocr" not in f
    assert f["segments"][0]["ref"] == "page 1 (OCR)"
    assert "Loyer 950" in gf.formater_pour_prompt([f])


def test_pdf_protege_par_mot_de_passe():
    f = gf.traiter_fichier(_Upload("secret.pdf", _pdf_chiffre()))
    assert f["erreur"] and "mot de passe" in f["erreur"]


# ── 5. Word ─────────────────────────────────────────────────────────────────

def test_docx_titres_listes_tableaux_entetes():
    f = gf.traiter_fichier(_Upload("rapport.docx", _docx()))
    assert f["erreur"] is None
    texte = f["contenu"]
    assert "# Rapport annuel" in texte and "## Risques" in texte
    assert "- Retard fournisseur" in texte
    assert "| Marge | 18 % |" in texte
    assert "Societe Alpha" in texte
    refs = [s["ref"] for s in f["segments"]]
    assert "section « Risques »" in refs and "en-têtes et pieds de page" in refs


# ── 6-7. Excel et CSV ───────────────────────────────────────────────────────

def test_xlsx_feuilles_statistiques_et_formules():
    f = gf.traiter_fichier(_Upload("ventes.xlsx", _xlsx()))
    assert f["erreur"] is None and f["type"] == "tableur"
    assert "Ventes" in f["resume"] and "Notes" in f["resume"]
    stats = next(s["texte"] for s in f["segments"] if "Ventes" in s["ref"] and "statistiques" in s["ref"])
    assert "CALCULÉES PAR L'APPLICATION" in stats
    assert "somme=13400" in stats and "max=9900" in stats
    assert "valeurs atypiques" in stats  # 9900 détecté comme anomalie
    # openpyxl n'enregistre pas de valeur calculée : la limite est signalée, pas inventée
    assert any("sans valeur en cache" in a for a in f["avertissements"])
    lignes = next(s for s in f["segments"] if "lignes 2–5" in s["ref"])
    assert "| 5 | Avril | Est | 9900 |" in lignes["texte"]


def test_csv_separateur_point_virgule():
    data = "produit;quantite;prix\nVis;100;0,05\nEcrou;250;0,03\nBoulon;80;0,12\n".encode("utf-8")
    f = gf.traiter_fichier(_Upload("stock.csv", data))
    assert f["erreur"] is None
    assert "point-virgule" in f["resume"]
    stats = next(s["texte"] for s in f["segments"] if "statistiques" in s["ref"])
    assert "quantite : n=3, somme=430" in stats
    assert "prix : n=3, somme=0.2," in stats  # décimales à la française reconnues
    assert "Ecrou" in gf.formater_pour_prompt([f])


def test_pptx_et_formats_texte():
    f = gf.traiter_fichier(_Upload("plan.pptx", _pptx()))
    assert f["erreur"] is None and f["segments"][0]["ref"] == "diapositive 1"
    assert "Lancement V2" in f["contenu"] and "Insister sur la sécurité" in f["contenu"]

    html = b"<html><script>alert(1)</script><h1>Titre</h1><p>Paragraphe utile</p></html>"
    h = gf.traiter_fichier(_Upload("page.html", html))
    assert "Paragraphe utile" in h["contenu"] and "alert" not in h["contenu"]

    j = gf.traiter_fichier(_Upload("conf.json", b'{"a": 1, "b": [1, 2]}'))
    assert '"b": [' in j["contenu"]

    x = gf.traiter_fichier(_Upload("bombe.xml", b'<!DOCTYPE x [<!ENTITY e "a">]><x>&e;</x>'))
    assert x["erreur"] is None and any("entités" in a for a in x["avertissements"])


# ── 8. Plusieurs pièces jointes, budget et documents volumineux ─────────────

def test_plusieurs_pieces_jointes_dans_un_message():
    fichiers = [
        gf.traiter_fichier(_Upload("rapport.pdf", _pdf_texte(["Anomalie A : stock negatif."]))),
        gf.traiter_fichier(_Upload("ventes.xlsx", _xlsx())),
        gf.traiter_fichier(_Upload("capture.png", _png_texte(["Erreur 504"]))),
    ]
    assert all(f["erreur"] is None for f in fichiers)
    gf.preparer_contenus_multimodaux(fichiers, question="Identifie les anomalies",
                                     llm_client=_FauxVision("Message : Erreur 504"))
    prompt = gf.formater_pour_prompt(fichiers, question="Identifie les anomalies")
    for attendu in ("rapport.pdf", "ventes.xlsx", "capture.png", "stock negatif", "9900", "Erreur 504"):
        assert attendu in prompt
    assert len(prompt) < config.FICHIERS_MAX_CHARS_PROMPT + 2500


def test_document_volumineux_selection_ciblee_avec_references():
    pages = [f"Page de remplissage numero {i}. " + "Texte sans rapport. " * 60 for i in range(1, 41)]
    pages[27] = "Clause de resiliation : preavis de trois mois par lettre recommandee. " * 3
    f = gf.traiter_fichier(_Upload("contrat.pdf", _pdf_texte(pages)))
    assert len(f["segments"]) == 40
    prompt = gf.formater_pour_prompt([f], question="Quel est le préavis de résiliation ?", max_chars=6000)
    assert "EXTRAITS sélectionnés" in prompt
    assert "--- page 28" in prompt and "preavis de trois mois" in prompt
    assert len(prompt) < 9000


def test_orchestrateur_transmet_le_contenu_au_llm(tmp_path, monkeypatch):
    from core.journal_amelioration import JournalAmelioration
    from core.memoire_longue import MemoireLongue
    from core.orchestrateur import OrchestrateurAssistant

    class _Retrieval:
        def rechercher(self, *a, **k):
            return []

        def doit_sabstenir(self, documents):
            return False

    class _LLM(_FauxVision):
        temperature = 0.3

        def __init__(self):
            super().__init__("Carré rouge, texte FACTURE 4271")
            self.messages = None

        def stream(self, messages, usage_holder=None):
            self.messages = messages
            yield "ok"

        def invoke(self, messages, **_):
            self.messages = messages
            return "ok", 1

    monkeypatch.setenv("DESACTIVER_WEB_SEARCH", "true")
    llm = _LLM()
    orch = OrchestrateurAssistant(retrieval=_Retrieval(), llm_client=llm)
    orch.memoire_longue = MemoireLongue(chemin=tmp_path / "m.db")
    orch.journal = JournalAmelioration(chemin=tmp_path / "j.jsonl")
    fichiers = [gf.traiter_fichier(_Upload("ventes.xlsx", _xlsx())),
                gf.traiter_fichier(_Upload("facture.png", _png_texte(["FACTURE 4271"])))]
    res = orch.repondre("Quel est le total des montants ?", fichiers=fichiers,
                        utiliser_corpus=False, stream=True)
    if res.stream is not None:
        "".join(res.stream)
    prompt = llm.messages[-1]["content"]
    assert "somme=13400" in prompt and "FACTURE 4271" in prompt
    assert len(llm.appels) == 1  # l'image est passée par la vision

    # « Fichier + corpus » avec une base vectorielle injoignable : la réponse reste possible
    class _QdrantEteint(_Retrieval):
        def rechercher(self, *a, **k):
            raise ConnectionRefusedError("[WinError 10061]")

    orch.retrieval = _QdrantEteint()
    res = orch.repondre("Quel est le total des montants ?", fichiers=fichiers[:1],
                        utiliser_corpus=True, stream=True)
    if res.stream is not None:
        "".join(res.stream)
    assert not res.abstention and "rag_indisponible" in res.analyse.notes
    assert "somme=13400" in llm.messages[-1]["content"]

    # Question sur la pièce jointe : pas de recherche web parasite, sauf demande explicite
    class _Web:
        def __init__(self):
            self.requetes = []

        def rechercher_multi(self, question):
            self.requetes.append(question)
            return []

    orch.retrieval = _Retrieval()
    orch.searcher = _Web()
    for question in ("Quel mois a la plus forte valeur dans ce graphique ?",
                     "Explique le fonctionnement technique de ce tableau"):
        res = orch.repondre(question, fichiers=fichiers[:1], utiliser_corpus=True, stream=True)
        if res.stream is not None:
            "".join(res.stream)
    assert orch.searcher.requetes == []
    res = orch.repondre("Quel est le prix actuel de ce produit en 2026 ?", fichiers=fichiers[:1],
                        utiliser_corpus=True, stream=True)
    if res.stream is not None:
        "".join(res.stream)
    assert len(orch.searcher.requetes) == 1


def test_erreurs_llm_sans_detail_technique():
    from generation.llm_client import LLMClient

    class APIConnectionError(Exception):
        pass

    erreurs = [
        APIConnectionError("Connection error."),
        Exception("Error code: 401 - {'error': {'message': 'Invalid API Key gsk_abc123'}}"),
        Exception("Error code: 400 - {'error': {'type': 'invalid_request_error'}}"),
    ]
    for exc in erreurs:
        llm = LLMClient.__new__(LLMClient)
        llm.modele, llm.temperature, llm.top_p = "m", 0.3, 1.0
        llm.presence_penalty = llm.frequency_penalty = 0.0
        llm.max_tokens = 100

        class _Completions:
            def create(self, **_):
                raise exc

        llm.client = type("C", (), {"chat": type("Ch", (), {"completions": _Completions()})()})()
        for texte in (llm.invoke([{"role": "user", "content": "x"}])[0],
                      "".join(llm.stream([{"role": "user", "content": "x"}]))):
            assert texte.startswith("❌ Erreur")
            assert "gsk_" not in texte and "{" not in texte and "Error code" not in texte


# ── 9-10. Retrait avant envoi / annulation du sélecteur ─────────────────────

def test_aucune_piece_jointe_aucun_contexte():
    # Fichier retiré avant l'envoi ou sélecteur annulé : chat_input ne transmet aucun fichier
    assert gf.formater_pour_prompt([]) == ""
    assert gf.preparer_contenus_multimodaux([], llm_client=_FauxVision()) == []


# ── 11-13. Formats interdits, taille, fichiers corrompus ou trompeurs ───────

@pytest.mark.parametrize("nom,attendu", [
    ("programme.exe", "non pris en charge"),
    ("ancien.doc", ".docx"),
    ("macro.xlsm", "macros"),
    ("dessin.svg", "PNG"),
    ("sans_extension", "non pris en charge"),
])
def test_formats_interdits(nom, attendu):
    f = gf.traiter_fichier(_Upload(nom, b"MZ\x90\x00binaire"))
    assert f["erreur"] and attendu in f["erreur"]


def test_fichier_trop_volumineux(monkeypatch):
    monkeypatch.setattr(config, "UPLOAD_MAX_BYTES", 1024)
    f = gf.traiter_fichier(_Upload("gros.txt", b"a" * 2048))
    assert f["erreur"] and "trop volumineux" in f["erreur"]


@pytest.mark.parametrize("nom,data,attendu", [
    ("faux.pdf", b"MZ\x90\x00 pas un pdf", "PDF"),
    ("tronque.pdf", b"%PDF-1.4\n1 0 obj << /Type /Catalog", "PDF"),
    ("virus.png", b"MZ\x90\x00\x03\x00\x00\x00", "image"),
    ("corrompu.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 40, "image"),
    ("faux.docx", b"PK\x03\x04 pas vraiment un zip", "corrompu"),
    ("vide.txt", b"", "vide"),
    ("binaire.txt", b"abc\x00\x01\x02", "binaires"),
])
def test_fichiers_corrompus_ou_trompeurs(nom, data, attendu):
    f = gf.traiter_fichier(_Upload(nom, data))
    assert f["erreur"] and attendu in f["erreur"]
    assert "Traceback" not in f["erreur"] and "Error" not in f["erreur"]


def test_xlsx_deguise_en_docx_et_bombe_zip(monkeypatch):
    f = gf.traiter_fichier(_Upload("rapport.docx", _xlsx()))
    assert f["erreur"] and "extension trompeuse" in f["erreur"]
    monkeypatch.setattr(config, "ARCHIVE_MAX_DECOMPRESSE", 1000)
    g = gf.traiter_fichier(_Upload("rapport.docx", _docx()))
    assert g["erreur"] and "décompressé" in g["erreur"]


def test_nom_de_fichier_assaini():
    f = gf.traiter_fichier(_Upload("../../etc/passwd.txt", b"root"))
    assert f["nom"] == "passwd.txt"


# ── 14-15. Isolation et persistance des références ──────────────────────────

def _message_avec_pieces(tmp_path):
    img = gf.traiter_fichier(_Upload("facture.png", _png_texte(["FACTURE 4271"])))
    gf.preparer_contenus_multimodaux([img], llm_client=_FauxVision())
    pdf = gf.traiter_fichier(_Upload("rapport.pdf", _pdf_texte(["Budget 48 000 euros"])))
    return [{"role": "user", "content": "Analyse", "fichiers": [img, pdf]},
            {"role": "assistant", "content": "Réponse"}]


def test_pieces_jointes_isolees_entre_utilisateurs(tmp_path):
    histo = HistoriqueSQLite(tmp_path / "h.db")
    cid = histo.creer_conversation("A", owner_id="u-a")
    histo.sauvegarder_conversation(cid, "A", _message_avec_pieces(tmp_path), owner_id="u-a")
    with pytest.raises(AccesRefuse):
        histo.charger_messages(cid, owner_id="u-b")
    assert histo.lister_conversations(owner_id="u-b") == []


def test_persistance_sans_binaire_et_question_de_suivi(tmp_path):
    histo = HistoriqueSQLite(tmp_path / "h.db")
    cid = histo.creer_conversation("A", owner_id="u-a")
    histo.sauvegarder_conversation(cid, "A", _message_avec_pieces(tmp_path), owner_id="u-a")
    msgs = histo.charger_messages(cid, owner_id="u-a")
    img, pdf = msgs[0]["fichiers"]
    for f in (img, pdf):
        assert not any(k in f for k in gf.CLES_TRANSITOIRES)
    assert img["image_omise"] and img["vision_ok"]
    assert "FACTURE 4271" in img["segments"][0]["texte"]
    assert pdf["segments"][0]["ref"] == "page 1"

    # Rechargement puis question de suivi sans pièce jointe : les documents restent en contexte
    msgs.append({"role": "user", "content": "Et le budget ?"})
    rappel = gf.fichiers_de_rappel(msgs)
    assert [f["nom"] for f in rappel] == ["facture.png", "rapport.pdf"]
    assert all(f["rappel"] for f in rappel) and "rappel" not in msgs[0]["fichiers"][0]
    prompt = gf.formater_pour_prompt(rappel, question="Et le budget ?")
    assert "joint plus tôt" in prompt and "48 000 euros" in prompt and "FACTURE 4271" in prompt


# ── Tests réels (modèle de vision et LLM Groq) ──────────────────────────────

@live
def test_live_vision_decrit_une_image():
    from generation.llm_client import LLMClient

    f = gf.traiter_fichier(_Upload("facture.png", _png_texte(["FACTURE 4271", "Total 350 EUR"])))
    gf.preparer_contenus_multimodaux([f], question="Décris cette image", llm_client=LLMClient())
    assert f["vision_ok"], f["avertissements"]
    assert "4271" in f["contenu"] and "rouge" in f["contenu"].lower()


@live
def test_live_capture_ecran_explication():
    from generation.llm_client import LLMClient

    png = _png_texte(["Connexion au serveur", "Erreur 504 : delai depasse", "[ Reessayer ]"],
                     taille=(700, 220), rouge=False)
    f = gf.traiter_fichier(_Upload("ecran.png", png))
    gf.preparer_contenus_multimodaux([f], question="Que signifie le message d'erreur ?",
                                     llm_client=LLMClient())
    assert f["vision_ok"] and "504" in f["contenu"]


@live
def test_live_ocr_pdf_numerise():
    from generation.llm_client import LLMClient

    f = gf.traiter_fichier(_Upload("scan.pdf", _pdf_scanne(["CONTRAT DE BAIL", "Loyer 950 euros"])))
    gf.preparer_contenus_multimodaux([f], question="Quel est le loyer ?", llm_client=LLMClient())
    assert f["segments"] and f["segments"][0]["ref"] == "page 1 (OCR)"
    assert "950" in f["contenu"]


def _orchestrateur_reel(tmp_path, monkeypatch):
    from core.journal_amelioration import JournalAmelioration
    from core.memoire_longue import MemoireLongue
    from core.orchestrateur import OrchestrateurAssistant

    class _Retrieval:
        def rechercher(self, *a, **k):
            return []

        def doit_sabstenir(self, documents):
            return False

    monkeypatch.setenv("DESACTIVER_WEB_SEARCH", "true")
    orch = OrchestrateurAssistant(retrieval=_Retrieval())
    orch.memoire_longue = MemoireLongue(chemin=tmp_path / "m.db")
    orch.journal = JournalAmelioration(chemin=tmp_path / "j.jsonl")
    return orch


def _repondre(orch, question, fichiers):
    res = orch.repondre(question, fichiers=fichiers, utiliser_corpus=False, stream=False)
    return (res.reponse_seule or res.reponse or "")


@live
def test_live_bout_en_bout_excel_word_pdf(tmp_path, monkeypatch):
    orch = _orchestrateur_reel(tmp_path, monkeypatch)
    rep = _repondre(orch, "Quel est le montant total des ventes et quel mois est atypique ?",
                    [gf.traiter_fichier(_Upload("ventes.xlsx", _xlsx()))])
    assert "13" in rep and "400" in rep and "Avril" in rep

    rep = _repondre(orch, "Quelle est la marge indiquée dans le document ?",
                    [gf.traiter_fichier(_Upload("rapport.docx", _docx()))])
    assert "18" in rep

    pdf = _pdf_texte(["Rapport Atlas. Introduction.", "Le budget total s'eleve a 48 000 euros.",
                      "Conclusion : livraison en avance."])
    rep = _repondre(orch, "Résume ce rapport en trois points.",
                    [gf.traiter_fichier(_Upload("rapport.pdf", pdf))])
    assert "48" in rep
