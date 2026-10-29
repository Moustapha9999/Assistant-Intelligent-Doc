"""Extraction structurée du contenu des pièces jointes.

Chaque extracteur renvoie des segments `{"ref", "texte"}` : la référence (page, feuille,
section, diapositive…) accompagne le texte jusqu'au prompt pour que les réponses soient
vérifiables. Les erreurs bloquantes lèvent `FichierRefuse` (message utilisateur).
"""

from __future__ import annotations

import base64
import csv
import io
import json
import re
import time
from html.parser import HTMLParser
from typing import Any, Dict, List, Optional

import config
from app.securite_fichiers import FichierRefuse


def _resultat(segments=None, resume="", avertissements=None, **extras) -> Dict[str, Any]:
    out = {
        "segments": [s for s in (segments or []) if (s.get("texte") or "").strip()],
        "resume": resume,
        "avertissements": list(avertissements or []),
    }
    out.update(extras)
    return out


def decoder_texte(contenu: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8"):
        try:
            return contenu.decode(enc)
        except UnicodeDecodeError:
            continue
    return contenu.decode("cp1252", errors="replace")


def _cellule(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace("\n", " ").replace("|", "/").strip()


def _tableau_markdown(lignes: List[List[Any]]) -> str:
    lignes = [[_cellule(c) for c in l] for l in lignes if l and any(_cellule(c) for c in l)]
    if not lignes:
        return ""
    n = max(len(l) for l in lignes)
    lignes = [l + [""] * (n - len(l)) for l in lignes]
    out = ["| " + " | ".join(lignes[0]) + " |", "|" + "---|" * n]
    out += ["| " + " | ".join(l) + " |" for l in lignes[1:]]
    return "\n".join(out)


def _image_page_b64(page, resolution: int = 130) -> str:
    rendu = page.to_image(resolution=resolution).original.convert("RGB")
    if max(rendu.size) > 2000:
        rendu.thumbnail((2000, 2000))
    buf = io.BytesIO()
    rendu.save(buf, format="JPEG", quality=85)
    return base64.b64encode(buf.getvalue()).decode("ascii")


# ── PDF ───────────────────────────────────────────────────────────────────────

def extraire_pdf(contenu: bytes) -> Dict[str, Any]:
    import pdfplumber

    try:
        pdf = pdfplumber.open(io.BytesIO(contenu))
        nb_pages = len(pdf.pages)
    except Exception as exc:
        txt = f"{exc!r} {exc.__context__!r}".lower()
        if "password" in txt or "encrypt" in txt:
            raise FichierRefuse(
                "Ce PDF est protégé par un mot de passe : retirez la protection puis réessayez."
            ) from exc
        raise FichierRefuse("Ce PDF est corrompu ou illisible.") from exc

    segments: List[Dict[str, str]] = []
    avert: List[str] = []
    pages_ocr: List[Dict[str, Any]] = []
    pages_figures: List[Dict[str, Any]] = []
    sans_texte: List[int] = []
    nb_tableaux = 0
    debut = time.monotonic()
    with pdf:
        limite = min(nb_pages, config.PDF_MAX_PAGES)
        for i in range(limite):
            if time.monotonic() - debut > config.FICHIERS_TIMEOUT_EXTRACTION_S:
                avert.append(
                    f"Extraction interrompue après {i} pages sur {nb_pages} (délai dépassé)."
                )
                break
            page = pdf.pages[i]
            num = i + 1
            try:
                texte = (page.extract_text() or "").strip()
                tableaux = page.extract_tables() or []
            except Exception:
                avert.append(f"Page {num} illisible : ignorée.")
                continue
            blocs = [texte] if texte else []
            for k, t in enumerate(tableaux, 1):
                md = _tableau_markdown(t)
                if md:
                    nb_tableaux += 1
                    blocs.append(f"Tableau {k} (page {num}) :\n{md}")
            if blocs:
                segments.append({"ref": f"page {num}", "texte": "\n\n".join(blocs)})
            a_images = bool(page.images)
            if len(texte) < 20:
                sans_texte.append(num)
                if len(pages_ocr) < config.PDF_OCR_MAX_PAGES:
                    try:
                        pages_ocr.append({"page": num, "image_b64": _image_page_b64(page, 150),
                                          "media_type": "image/jpeg"})
                    except Exception:
                        pass
            elif a_images and len(pages_figures) < config.PDF_FIGURES_MAX_PAGES:
                try:
                    pages_figures.append({"page": num, "image_b64": _image_page_b64(page),
                                          "media_type": "image/jpeg"})
                except Exception:
                    pass
        if nb_pages > limite:
            avert.append(f"Seules les {limite} premières pages sur {nb_pages} ont été lues.")

    if sans_texte:
        n_ocr = len(pages_ocr)
        avert.append(
            f"{len(sans_texte)} page(s) sans texte sélectionnable (document numérisé ?) : "
            f"OCR par le modèle de vision sur {n_ocr} page(s)"
            + (f", {len(sans_texte) - n_ocr} non traitée(s) (limite {config.PDF_OCR_MAX_PAGES})."
               if len(sans_texte) > n_ocr else ".")
        )
    if not segments and not pages_ocr:
        raise FichierRefuse("Aucun texte n'a pu être extrait de ce PDF.")
    resume = f"PDF · {nb_pages} page{'s' if nb_pages > 1 else ''}"
    if nb_tableaux:
        resume += f" · {nb_tableaux} tableau{'x' if nb_tableaux > 1 else ''}"
    return _resultat(segments, resume, avert, pages_ocr=pages_ocr, pages_figures=pages_figures)


# ── Word (.docx) ──────────────────────────────────────────────────────────────

def extraire_docx(contenu: bytes) -> Dict[str, Any]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(contenu))
    except Exception as exc:
        raise FichierRefuse("Ce document Word est corrompu ou illisible.") from exc

    segments: List[Dict[str, str]] = []
    section, lignes = "début du document", []
    nb_tableaux = 0

    def clore():
        if lignes:
            segments.append({"ref": f"section « {section} »", "texte": "\n".join(lignes)})

    for el in document.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, document)
            texte = (p.text or "").strip()
            if not texte:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"(?:heading|titre)\s*(\d)?", style, re.I)
            if m or style.lower() == "title":
                clore()
                section, lignes = texte[:80], []
                niveau = int(m.group(1)) if m and m.group(1) else 1
                lignes.append("#" * min(niveau, 6) + " " + texte)
            elif "list" in style.lower() or p._p.pPr is not None and p._p.pPr.numPr is not None:
                lignes.append(f"- {texte}")
            else:
                lignes.append(texte)
        elif tag == "tbl":
            nb_tableaux += 1
            t = Table(el, document)
            rows = [[c.text for c in r.cells] for r in t.rows]
            md = _tableau_markdown(rows)
            if md:
                lignes.append(f"Tableau {nb_tableaux} :\n{md}")
    clore()

    entetes = []
    for s in document.sections:
        for part in (s.header, s.footer):
            for p in part.paragraphs:
                t = (p.text or "").strip()
                if t and t not in entetes:
                    entetes.append(t)
    if entetes:
        segments.append({"ref": "en-têtes et pieds de page", "texte": "\n".join(entetes)})

    avert = []
    nb_images = len(document.inline_shapes)
    if nb_images:
        avert.append(f"{nb_images} image(s) intégrée(s) au document non analysée(s).")
    if not segments:
        raise FichierRefuse("Ce document Word est vide ou ne contient aucun texte extractible.")
    resume = f"Word · {len(segments)} section(s)"
    if nb_tableaux:
        resume += f" · {nb_tableaux} tableau{'x' if nb_tableaux > 1 else ''}"
    return _resultat(segments, resume, avert)


# ── Tableurs (xlsx, xls, csv) ─────────────────────────────────────────────────

def _entetes_uniques(brut: List[Any]) -> List[str]:
    vus: Dict[str, int] = {}
    out = []
    for i, h in enumerate(brut):
        nom = _cellule(h) or f"colonne_{i + 1}"
        if nom in vus:
            vus[nom] += 1
            nom = f"{nom}_{vus[nom]}"
        else:
            vus[nom] = 1
        out.append(nom)
    return out


def _nombre(v: Any) -> Optional[float]:
    """Valeur numérique d'une cellule (accepte « 1 234,5 », « 12 % », « 3.2 ») ; sinon None."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace("\u00a0", "").replace(" ", "").rstrip("%€$")
    if re.fullmatch(r"-?\d+(,\d+)?", t):
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def _quantile(valeurs: List[float], q: float) -> float:
    pos = (len(valeurs) - 1) * q
    bas = int(pos)
    haut = min(bas + 1, len(valeurs) - 1)
    return valeurs[bas] + (valeurs[haut] - valeurs[bas]) * (pos - bas)


def _segments_feuille(nom: str, lignes: List[List[Any]], tronquee: bool,
                      lignes_par_bloc: int = 40) -> List[Dict[str, str]]:
    """Structure + statistiques calculées + données par blocs (n° de ligne Excel conservés)."""
    import statistics
    from collections import Counter

    lignes = [list(l) for l in lignes if l is not None]
    idx = [i for i, l in enumerate(lignes) if any(_cellule(c) for c in l)]
    if not idx:
        return []
    debut = idx[0]
    entetes = _entetes_uniques(lignes[debut])
    largeur = len(entetes)
    corps = [
        l[:largeur] + [None] * (largeur - len(l))
        for l in lignes[debut + 1:]
        if any(_cellule(c) for c in l[:largeur])
    ]
    colonnes = {h: [r[j] for r in corps] for j, h in enumerate(entetes)}
    ref = f"feuille « {nom} »"

    numeriques: Dict[str, List[float]] = {}
    for col, vals in colonnes.items():
        remplies = [v for v in vals if _cellule(v)]
        nums = [n for n in (_nombre(v) for v in remplies) if n is not None]
        if nums and len(nums) >= 0.8 * len(remplies):
            numeriques[col] = nums

    struct = [
        f"Feuille : {nom}",
        f"En-têtes (ligne {debut + 1}) : {', '.join(entetes)}",
        f"Lignes de données lues : {len(corps)}" + (" (feuille tronquée)" if tronquee else ""),
        f"Colonnes numériques : {', '.join(numeriques) or 'aucune'}",
    ]
    stats = [
        "Statistiques CALCULÉES PAR L'APPLICATION à partir des valeurs lues — "
        "ce ne sont pas des valeurs présentes dans le fichier :"
    ]
    for col, s in numeriques.items():
        tri = sorted(s)
        q1, q3 = _quantile(tri, 0.25), _quantile(tri, 0.75)
        iqr = q3 - q1
        aberr = sum(1 for x in s if x < q1 - 1.5 * iqr or x > q3 + 1.5 * iqr) if iqr else 0
        ecart = statistics.stdev(s) if len(s) > 1 else 0.0
        stats.append(
            f"- {col} : n={len(s)}, somme={sum(s):.6g}, moyenne={statistics.fmean(s):.6g}, "
            f"médiane={statistics.median(s):.6g}, min={tri[0]:.6g}, max={tri[-1]:.6g}, "
            f"écart-type={ecart:.6g}" + (f", valeurs atypiques (IQR)={aberr}" if aberr else "")
        )
    for col, vals in colonnes.items():
        if col in numeriques:
            continue
        textes = [_cellule(v) for v in vals if _cellule(v)]
        if textes:
            top = Counter(textes).most_common(5)
            stats.append(
                f"- {col} (texte) : {len(set(textes))} valeurs distinctes, "
                f"{len(vals) - len(textes)} vide(s), fréquentes : "
                + ", ".join(f"{k} ({v})" for k, v in top)
            )
    segs = [{"ref": f"{ref} — structure", "texte": "\n".join(struct)},
            {"ref": f"{ref} — statistiques calculées", "texte": "\n".join(stats)}]

    lignes_txt = []
    for i, l in enumerate(lignes[debut + 1:], start=debut + 2):
        if any(_cellule(c) for c in l[:largeur]):
            lignes_txt.append((i, [_cellule(c) for c in l[:largeur]]))
    for b in range(0, len(lignes_txt), lignes_par_bloc):
        bloc = lignes_txt[b:b + lignes_par_bloc]
        md = _tableau_markdown([["ligne"] + entetes] + [[str(n)] + v for n, v in bloc])
        segs.append({"ref": f"{ref}, lignes {bloc[0][0]}–{bloc[-1][0]}", "texte": md})
    return segs


def extraire_xlsx(contenu: bytes) -> Dict[str, Any]:
    import openpyxl

    try:
        wb_val = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
        wb_form = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=False)
    except Exception as exc:
        raise FichierRefuse("Ce classeur Excel est corrompu ou illisible.") from exc

    segments, avert = [], []
    noms = wb_val.sheetnames
    for nom in noms:
        ws_v, ws_f = wb_val[nom], wb_form[nom]
        if getattr(ws_v, "sheet_state", "visible") != "visible":
            avert.append(f"Feuille masquée « {nom} » lue également.")
        lignes, formules, sans_valeur = [], 0, 0
        tronquee = False
        for n, (rv, rf) in enumerate(zip(ws_v.iter_rows(values_only=True),
                                         ws_f.iter_rows(values_only=True))):
            if n > config.TABLEUR_MAX_LIGNES:
                tronquee = True
                break
            for v, f in zip(rv, rf):
                if isinstance(f, str) and f.startswith("="):
                    formules += 1
                    if v is None:
                        sans_valeur += 1
            lignes.append(list(rv))
        if tronquee:
            avert.append(f"Feuille « {nom} » tronquée à {config.TABLEUR_MAX_LIGNES} lignes.")
        if formules:
            msg = f"Feuille « {nom} » : {formules} formule(s) ; valeurs calculées lues depuis le cache Excel."
            if sans_valeur:
                msg += (f" {sans_valeur} formule(s) sans valeur en cache (classeur jamais recalculé) : "
                        "ces cellules apparaissent vides.")
            avert.append(msg)
        segments += _segments_feuille(nom, lignes, tronquee)
    wb_val.close()
    wb_form.close()
    if not segments:
        raise FichierRefuse("Ce classeur ne contient aucune donnée.")
    return _resultat(segments, f"Excel · {len(noms)} feuille(s) : {', '.join(noms)[:120]}", avert)


def extraire_xls(contenu: bytes) -> Dict[str, Any]:
    import xlrd

    try:
        wb = xlrd.open_workbook(file_contents=contenu, on_demand=True)
    except Exception as exc:
        raise FichierRefuse("Ce classeur Excel (.xls) est corrompu ou illisible.") from exc
    segments, avert = [], ["Format .xls : formules non disponibles, seules les valeurs enregistrées sont lues."]
    noms = wb.sheet_names()
    for nom in noms:
        sh = wb.sheet_by_name(nom)
        n = min(sh.nrows, config.TABLEUR_MAX_LIGNES + 1)
        tronquee = sh.nrows > n
        if tronquee:
            avert.append(f"Feuille « {nom} » tronquée à {config.TABLEUR_MAX_LIGNES} lignes.")
        segments += _segments_feuille(nom, [sh.row_values(i) for i in range(n)], tronquee)
    wb.release_resources()
    if not segments:
        raise FichierRefuse("Ce classeur ne contient aucune donnée.")
    return _resultat(segments, f"Excel 97-2003 · {len(noms)} feuille(s)", avert)


def extraire_csv(contenu: bytes, tsv: bool = False) -> Dict[str, Any]:
    texte = decoder_texte(contenu)
    if not texte.strip():
        raise FichierRefuse("Le fichier CSV est vide.")
    try:
        delim = "\t" if tsv else csv.Sniffer().sniff(texte[:20000], delimiters=",;\t|").delimiter
    except csv.Error:
        delim = ","
    lecteur = csv.reader(io.StringIO(texte), delimiter=delim)
    lignes, tronquee = [], False
    try:
        for i, l in enumerate(lecteur):
            if i > config.TABLEUR_MAX_LIGNES:
                tronquee = True
                break
            lignes.append(l)
    except csv.Error as exc:
        raise FichierRefuse("Le fichier CSV est mal formé.") from exc
    segments = _segments_feuille("CSV", lignes, tronquee)
    avert = [f"Fichier tronqué à {config.TABLEUR_MAX_LIGNES} lignes."] if tronquee else []
    if not segments:
        raise FichierRefuse("Le fichier CSV ne contient aucune donnée.")
    sep = {"\t": "tabulation", ",": "virgule", ";": "point-virgule", "|": "barre"}.get(delim, delim)
    return _resultat(segments, f"CSV · {len(lignes)} ligne(s) · séparateur {sep}", avert)


# ── PowerPoint (.pptx) ────────────────────────────────────────────────────────

def extraire_pptx(contenu: bytes) -> Dict[str, Any]:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(contenu))
    except Exception as exc:
        raise FichierRefuse("Cette présentation est corrompue ou illisible.") from exc
    segments, nb_images = [], 0
    for n, slide in enumerate(prs.slides, 1):
        blocs = []
        for shape in slide.shapes:
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                nb_images += 1
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                md = _tableau_markdown(rows)
                if md:
                    blocs.append(md)
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    blocs.append(t)
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                blocs.append(f"Notes de l'orateur : {notes}")
        if blocs:
            segments.append({"ref": f"diapositive {n}", "texte": "\n".join(blocs)})
    avert = [f"{nb_images} image(s) des diapositives non analysée(s)."] if nb_images else []
    if not segments:
        raise FichierRefuse("Cette présentation ne contient aucun texte extractible.")
    return _resultat(segments, f"PowerPoint · {len(prs.slides)} diapositive(s)", avert)


# ── Texte, code, HTML, JSON, XML ──────────────────────────────────────────────

class _TexteHTML(HTMLParser):
    _IGNORES = {"script", "style", "noscript", "template", "svg"}
    _BLOCS = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.morceaux: List[str] = []
        self._ignore = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._IGNORES:
            self._ignore += 1
        elif tag in self._BLOCS:
            self.morceaux.append("\n")
        if tag in {"h1", "h2", "h3"}:
            self.morceaux.append("#" * int(tag[1]) + " ")

    def handle_endtag(self, tag):
        if tag in self._IGNORES and self._ignore:
            self._ignore -= 1

    def handle_data(self, data):
        if not self._ignore:
            self.morceaux.append(data)


def _segmenter_texte(texte: str, prefixe: str, taille: int = 3000) -> List[Dict[str, str]]:
    lignes = texte.splitlines()
    segs, courant, debut = [], [], 1
    for i, l in enumerate(lignes, 1):
        courant.append(l)
        if sum(len(x) + 1 for x in courant) >= taille:
            segs.append({"ref": f"{prefixe} lignes {debut}–{i}".strip(), "texte": "\n".join(courant)})
            courant, debut = [], i + 1
    if courant:
        ref = prefixe if not segs else f"{prefixe} lignes {debut}–{len(lignes)}"
        segs.append({"ref": ref.strip() or "contenu intégral", "texte": "\n".join(courant)})
    return segs


def extraire_texte_brut(contenu: bytes, ext: str) -> Dict[str, Any]:
    texte = decoder_texte(contenu)
    avert = []
    if ext in {".html", ".htm"}:
        p = _TexteHTML()
        p.feed(texte)
        texte = re.sub(r"\n\s*\n+", "\n\n", "".join(p.morceaux)).strip()
        avert.append("Page HTML : seul le texte visible est extrait (scripts et styles ignorés).")
    elif ext == ".json":
        try:
            texte = json.dumps(json.loads(texte), ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            avert.append("JSON invalide : contenu transmis tel quel.")
    elif ext == ".xml":
        try:
            from defusedxml import ElementTree as DET
            DET.fromstring(texte.encode("utf-8"))
        except ImportError:
            pass
        except Exception:
            avert.append("XML mal formé ou contenant des entités interdites : contenu transmis comme texte.")
    if not texte.strip():
        raise FichierRefuse("Le fichier est vide.")
    nb = texte.count("\n") + 1
    return _resultat(_segmenter_texte(texte, ""), f"{nb} ligne(s)", avert)


EXTRACTEURS = {
    ".pdf": extraire_pdf,
    ".docx": extraire_docx,
    ".xlsx": extraire_xlsx,
    ".xls": extraire_xls,
    ".csv": extraire_csv,
    ".tsv": lambda c: extraire_csv(c, tsv=True),
    ".pptx": extraire_pptx,
}


def extraire(contenu: bytes, ext: str) -> Optional[Dict[str, Any]]:
    fn = EXTRACTEURS.get(ext)
    if fn is not None:
        return fn(contenu)
    return extraire_texte_brut(contenu, ext)
